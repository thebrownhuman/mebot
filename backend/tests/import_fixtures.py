"""Files for the import tests, built in memory (no binary fixtures in the repo)."""

import io
import zipfile


def _esc(s: str) -> str:
    return s.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")


def make_pdf(lines: list[str]) -> bytes:
    """A one-page PDF with real text (Helvetica), readable by pypdf."""
    body = " ".join(f"({_esc(line)}) Tj T*" for line in lines)
    stream = f"BT /F1 10 Tf 40 780 Td 12 TL {body} ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    return out.getvalue()


def make_encrypted_pdf(lines: list[str]) -> bytes:
    from pypdf import PdfReader, PdfWriter

    w = PdfWriter(clone_from=PdfReader(io.BytesIO(make_pdf(lines))))
    w.encrypt("secret")
    out = io.BytesIO()
    w.write(out)
    return out.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    from docx import Document

    d = Document()
    for p in paragraphs:
        d.add_paragraph(p)
    if table:
        t = d.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                t.cell(r, c).text = cell
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


def make_docm(paragraphs: list[str]) -> bytes:
    """A DOCX whose content types declare a macro-enabled document."""
    src = zipfile.ZipFile(io.BytesIO(make_docx(paragraphs)))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            data = src.read(info.filename)
            if info.filename == "[Content_Types].xml":
                data = data.replace(
                    b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml",
                    b"application/vnd.ms-word.document.macroEnabled.main+xml",
                )
            z.writestr(info.filename, data)
    return out.getvalue()


def make_zip(files: dict[str, bytes | str], compression=zipfile.ZIP_DEFLATED) -> bytes:
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", compression) as z:
        for name, data in files.items():
            z.writestr(name, data.encode("utf-8") if isinstance(data, str) else data)
    return out.getvalue()


OLE_DOC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504
EXE = b"MZ\x90\x00" + b"\x00" * 200
